from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "BusBuddy_Project_Report_Updated.docx"


def insert_before(anchor, element):
    anchor._p.addprevious(element._element if hasattr(element, "_element") else element)


def insert_paragraph_before(anchor, text, style="Normal"):
    paragraph = anchor._parent.add_paragraph(text, style=style)
    insert_before(anchor, paragraph)
    return paragraph


def insert_table_before(anchor, document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = value
    insert_before(anchor, table)
    return table


def main():
    document = Document(REPORT)

    # Add the new subsection to the existing contents list.
    contents_anchor = next(
        p for p in document.paragraphs if p.text.strip() == "7. Architecture Diagram — To Be Designed"
    )
    contents_entry = contents_anchor._parent.add_paragraph(
        "6.1 Implementation Progress Update", style=contents_anchor.style
    )
    contents_anchor._p.addprevious(contents_entry._p)

    # Add a locally scoped progress update immediately before the reserved architecture section.
    section_anchor = next(
        p for p in document.paragraphs if p.text.strip() == "Architecture Diagram — To Be Designed"
    )
    insert_paragraph_before(section_anchor, "6.1 Implementation Progress Update", style="Heading 1")
    insert_paragraph_before(
        section_anchor,
        "Status date: 29 August 2026. The project has moved from requirements and UI exploration into an Android-first Flutter implementation for the Vellore / VIT demonstration context.",
    )
    insert_paragraph_before(
        section_anchor,
        "The current build deliberately implements a dependable core journey with local demonstration data. It does not yet claim live transit availability, production ticketing, AI action execution or field-validated bus recognition.",
    )

    progress_rows = [
        (
            "Completed",
            "Flutter Android-first project scaffold and accessible theme.",
            "Implementation exists in lib/; large targets, contrast and reusable semantic components are in place.",
        ),
        (
            "Completed",
            "Local transport data and repository layer.",
            "A 12-stop VIT Vellore to Katpadi Railway Station demonstration corridor supports deterministic route search.",
        ),
        (
            "Completed",
            "Core journey flow.",
            "Home → route search → route details → active journey, with explicit journey state and reactive updates.",
        ),
        (
            "Completed",
            "Home-screen redesign.",
            "Plan journey, start/destination controls, Depart now, stacked active-trip/tickets/safety cards and full-width Talk to BusBuddy action follow the accessibility review and Stitch reference direction.",
        ),
        (
            "Completed",
            "Accessibility behaviour in the prototype.",
            "TalkBack-oriented labels, focus order, live-region announcements, repeat/cancel recovery and uncertainty-aware status copy are represented in the core screens.",
        ),
        (
            "In progress",
            "Verification cleanup.",
            "Implementation review reports recorded passing Flutter tests and clean analysis; the workspace still contains a stale generated widget_test.dart that must be removed or updated before a fresh final run.",
        ),
        (
            "Planned",
            "Next implementation slice.",
            "Remove the stale test, rerun verification, then add map/OSM-backed route context and a deterministic simulated live-tracking state before external integrations.",
        ),
    ]
    insert_table_before(
        section_anchor,
        document,
        ("Status", "Workstream", "Current progress / evidence"),
        progress_rows,
    )
    insert_paragraph_before(
        section_anchor,
        "Not yet implemented: production authentication, GPS permissions, live vehicle feeds and ETA freshness, Gemini Live voice actions, computer-vision bus verification, UWB, ticketing/payment, emergency-contact sharing and driver communication. These remain staged requirements and require privacy, safety and field-validation work before deployment.",
    )
    insert_paragraph_before(
        section_anchor,
        "The immediate project milestone is therefore a clean, repeatable core-journey demo that can be evaluated with visually impaired participants, older users and typical commuters. Success should be measured through completion, errors, recovery actions, confidence and accessibility feedback rather than visual polish alone.",
    )

    document.save(REPORT)


if __name__ == "__main__":
    main()

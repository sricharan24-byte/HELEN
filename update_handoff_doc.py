from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

OUT = Path('/home/pavan/BusBuddy/BusBuddy_Project_Handoff_Document_Updated.docx')

def find(doc, starts):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts): return p
    raise ValueError(f'Paragraph not found: {starts}')

def set_text(p, text):
    for child in list(p._p):
        if child.tag.endswith('}r') or child.tag.endswith('}hyperlink'): p._p.remove(child)
    p.add_run(text)

def insert_after(anchor, text, style=None):
    new_p = OxmlElement('w:p'); anchor._p.addnext(new_p)
    p = Paragraph(new_p, anchor._parent)
    if style: p.style = style
    p.add_run(text); return p

def insert_before(anchor, text, style=None):
    new_p = OxmlElement('w:p'); anchor._p.addprevious(new_p)
    p = Paragraph(new_p, anchor._parent)
    if style: p.style = style
    p.add_run(text); return p

def replace_cell(table, row, col, text): table.cell(row, col).text = text

doc = Document(OUT)
if any(p.text.strip().startswith('16.2 Adaptive UI and User-Created Interface') for p in doc.paragraphs):
    print(f'{OUT} already contains the latest handoff updates')
    raise SystemExit(0)

set_text(find(doc, 'BusBuddy is a proposed accessibility-first public transport mobile application.'), 'BusBuddy is an accessibility-first, adaptive public transport assistant designed primarily for visually impaired users and older adults, while remaining usable by general commuters. It combines route planning, real-time bus tracking, accessible journey assistance, conversational AI, computer vision, and optional precision boarding assistance. The interface can adapt around frequently used features and journeys while allowing users to customize their own interface.')

users = doc.tables[1]
for row, label, need in [
    (0, 'Visually impaired users', 'Screen-reader compatibility, voice interaction, spoken journey state, accessible alerts, and user-controlled adaptive shortcuts'),
    (1, 'Older adults', 'Simple navigation, clear terminology, fewer confusing choices, larger touch targets, and predictable personalization'),
    (2, 'General commuters', 'Fast route search, bus status, ETA, clear journey information, and optional personalization'),
    (3, 'Low-vision users', 'Readable typography, scalable text, strong contrast, large controls, and accessible personalization controls'),
    (4, 'First-time/unfamiliar commuters', 'Clear route explanations and step-by-step journey guidance'),
]:
    replace_cell(users, row, 0, label); replace_cell(users, row, 1, need)

set_text(find(doc, 'Adaptive UI was intentionally deferred.'), 'Adaptive UI is now part of the project scope. The system may surface frequently used features, routes, destinations, and interaction methods as shortcuts, but users must be able to accept, reject, reset, or disable adaptive changes. Adaptation should assist the user, not take control away from the user.')

features = doc.tables[2]
for label, detail, priority in [
    ('Adaptive UI', 'Usage-based shortcuts for frequent destinations, routes, features, and interaction methods; user-controlled accept/reject/reset/disable', 'Core / Basic personalization'),
    ('User-Created Interface', 'Add, remove, rearrange, and reset home-screen components such as Voice Assistant, Route Search, Live Tracking, Favourite Routes, Journey Assistant, Map, Tickets, Notifications, and Emergency Contact; support drag-and-drop plus TalkBack, voice-command, and list-ordering alternatives', 'Advanced'),
    ('Ticketing & Payment', 'Future workflow: show ticket information and price, obtain explicit confirmation and authentication, then use a payment system', 'Advanced / Future'),
    ('AI-Assisted Driver Communication', 'Potential future call/message flow with confirmation, consent, privacy, authentication, driver contact availability, and operator-policy safeguards', 'Experimental / Future'),
]:
    row = features.add_row().cells; row[0].text, row[1].text, row[2].text = label, detail, priority

gemini = find(doc, 'Gemini Live should sit on top of the application')
set_text(gemini, 'Gemini Live acts as a conversational control layer over BusBuddy rather than merely a question-answering assistant. It may invoke authorized application functions such as route search, live-bus lookup, journey start, remaining-stop lookup, and favourite-route display, using structured app/backend data as the source of truth.')
insert_after(gemini, 'Users may say: “Find a bus from VIT to Katpadi,” “Where is my bus?”, “How many stops are left?”, “Start this journey,” or “Show me my favourite route.” Gemini must not guess transport facts or independently trigger sensitive actions.')

heading = insert_before(find(doc, '17. Accessibility Implementation Checklist'), '16.2 Adaptive UI and User-Created Interface', 'Heading 2')
insert_after(heading, 'Automatic usage-based adaptation may identify non-sensitive patterns such as frequently selected destinations, routes, features, and interaction methods. For example, repeated VIT → Katpadi searches may become a home-screen shortcut.')
insert_after(heading, 'Users must retain control: they can accept, reject, reset, or disable suggestions. A user-created interface may support adding, removing, and rearranging components. Drag-and-drop is appropriate for users who can use touch; TalkBack actions, voice commands, and accessible list-based ordering must provide equivalent alternatives.')
journey = insert_before(find(doc, '18. UX / Screen Inventory'), '17.1 Adaptive Journey', 'Heading 2')
insert_after(journey, 'User completes journeys → system observes non-sensitive usage patterns → frequent actions are identified → adaptive shortcuts are suggested → user accepts or rejects → personalized home screen')

tiers = doc.tables[12]
replace_cell(tiers, 1, 0, 'Core / Must Have'); replace_cell(tiers, 1, 1, 'Accessible Flutter UI, route search, route visualization, OpenStreetMap integration, bus stops, GPS/location, real-time or simulated bus tracking, ETA, Journey Assistant, TalkBack compatibility, TTS, voice interaction, accessible alerts, and basic personalization')
replace_cell(tiers, 2, 0, 'Advanced'); replace_cell(tiers, 2, 1, 'Gemini Live, conversational application control, computer vision, adaptive UI, user-created/customizable UI, and UWB precision boarding')
replace_cell(tiers, 3, 0, 'Future / Experimental'); replace_cell(tiers, 3, 1, 'AI-assisted driver communication, AI-mediated driver conversation, real-world transport-authority integration, real payment integration, advanced ML-based ETA, and large-scale deployment')

insert_after(find(doc, 'Do not assume every Android phone supports UWB'), 'Gemini must never independently initiate a financial transaction. Ticketing/payment requires explicit user confirmation and appropriate authentication. Driver communication, if explored later, requires user confirmation, consent, privacy safeguards, driver contact availability, telephony integration, authentication, and transport-operator policies; the AI should identify itself clearly to the driver.')

set_text(find(doc, 'I am building BusBuddy, a two-member semester project'), 'I am building BusBuddy, a two-member semester project for Usability Design of Software Applications. BusBuddy is an adaptive, accessibility-first Flutter public-transport assistant focused on Vellore City, Tamil Nadu, especially VIT Vellore and nearby locations/stops. It supports route/journey usability, OpenStreetMap-based maps and routing, real-time bus tracking and ETA, screen-reader-compatible and voice-first interaction, journey progress/stop alerts, usage-based shortcuts, user-controlled interface customization, and advanced Gemini Live conversational + computer-vision assistance. Gemini is an authorized control layer over app/backend data, not the source of transport facts. Ticketing/payment and driver communication are future/experimental concepts requiring explicit confirmation, authentication, consent, privacy, and operator safeguards. UWB near a bus entrance remains experimental and requires compatible hardware. Help me continue from this baseline with a simple two-person architecture, research-grounded decisions, and staged implementation.')
set_text(find(doc, 'This document captures the project'), 'This document captures the project\'s updated agreed direction and planning decisions. BusBuddy is adaptive and accessibility-first, prioritizing visually impaired users and older adults while remaining useful to general commuters. Adaptive UI and user-created customization are in scope with explicit user control. Gemini Live is a grounded conversational control layer. Ticketing/payment and AI-assisted driver communication remain advanced/future concepts. Geographic scope is fixed to Vellore City, with VIT Vellore and nearby stops as primary prototype points. Production Flutter implementation is still pending.')

doc.save(OUT)
print(OUT)
